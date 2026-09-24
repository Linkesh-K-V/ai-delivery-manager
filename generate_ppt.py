from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_title(slide, text):
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)
        return title_box

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide1, RGBColor(248, 250, 252))

    tbox = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "AI Delivery Manager"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(37, 99, 235)

    p2 = tf1.add_paragraph()
    p2.text = "An LLM-Powered Intelligent Workforce Management System for Dynamic Task Allocation & Workload Optimization"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(71, 85, 105)
    p2.space_after = Pt(40)

    p3 = tf1.add_paragraph()
    p3.text = "Presented by: LINKESH K V\nRoll No: CB.EN.U4ARE23023\nCourse: 25ARE471 Introduction to Generative AI"
    p3.font.size = Pt(18)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(15, 23, 42)

    # --- SLIDE 2: Motivation / Pain Points ---
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide2, RGBColor(255, 255, 255))
    add_title(slide2, "Motivation / Pain Points")

    cbox2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = cbox2.text_frame
    tf2.word_wrap = True

    points2 = [
        ("Cognitive Overload on Engineering Leads: ", "Managers manually juggle overlapping skills, deadlines, and fragmented calendar meetings across distributed teams."),
        ("Probabilistic LLM Failures: ", "Pure LLM architectures inherently hallucinate arithmetic—miscalculating working hours, deadlines, and schedule conflicts."),
        ("Opaque Decision-Making: ", "Traditional automated heuristic algorithms lack natural-language explainability, failing to communicate trade-offs and capacity risks."),
        ("Communication Latency: ", "Absence requests and task blockers are reported reactively through static forms, creating organizational lag and missed delivery deadlines.")
    ]
    for i, (title, desc) in enumerate(points2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(16)
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(18)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(18)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # --- SLIDE 3: Problem Definition & Objectives ---
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide3, RGBColor(255, 255, 255))
    add_title(slide3, "Problem Definition & Objectives")

    cbox3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = cbox3.text_frame
    tf3.word_wrap = True

    p_prob = tf3.paragraphs[0]
    p_prob.text = "Problem Definition:"
    p_prob.font.bold = True
    p_prob.font.size = Pt(20)
    p_prob.font.color.rgb = RGBColor(37, 99, 235)

    p_prob_desc = tf3.add_paragraph()
    p_prob_desc.text = "To engineer an enterprise task-allocation framework that eliminates LLM arithmetic hallucinations by decoupling language reasoning from deterministic business logic, retaining full human governance."
    p_prob_desc.font.size = Pt(18)
    p_prob_desc.font.color.rgb = RGBColor(51, 65, 85)
    p_prob_desc.space_after = Pt(20)

    p_obj = tf3.add_paragraph()
    p_obj.text = "Core Objectives:"
    p_obj.font.bold = True
    p_obj.font.size = Pt(20)
    p_obj.font.color.rgb = RGBColor(37, 99, 235)

    objs = [
        "1. Extract structured task requirements and estimated effort from natural language prompts using local LLM inference.",
        "2. Deterministically calculate real-time calendar capacity and weighted skill compatibility without model hallucination.",
        "3. Synthesize factual workload signals into explainable assignment trade-offs via a reasoning agent.",
        "4. Enforce strict Human-in-the-Loop (HITL) authority: AI recommends → Human approves → Backend executes.",
        "5. Enable Level-1 Bounded What-If Simulations to project absence impacts without production state mutation."
    ]
    for obj in objs:
        p = tf3.add_paragraph()
        p.text = obj
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(71, 85, 105)
        p.space_after = Pt(8)

    # --- SLIDE 4: Solution Formulation (Architecture Diagram) ---
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide4, RGBColor(255, 255, 255))
    add_title(slide4, "Solution Formulation – Architecture Diagram")

    # Draw Architecture Flow Blocks
    blocks = [
        ("Input Stage", "Natural Language Prompt\n(Manager / Employee)", Inches(0.8), Inches(2.2), Inches(2.6), Inches(1.8)),
        ("GenAI Parsing", "Local Llama 3 (8B) +\nInstructor Validation\n(Extracts Skills & Effort)", Inches(3.8), Inches(2.2), Inches(2.6), Inches(1.8)),
        ("Deterministic Core", "Workload & Calendar Engine\n+ Candidate Scorer\n(Zero LLM Math)", Inches(6.8), Inches(2.2), Inches(2.6), Inches(1.8)),
        ("Reasoning Layer", "LLM Reasoning Agent\n(Generates Tradeoffs &\nConfidence Level)", Inches(9.8), Inches(2.2), Inches(2.6), Inches(1.8)),
        ("Human Gatekeeper", "Manager Approval UI\n(Approve / Reject / Modify)", Inches(3.8), Inches(4.7), Inches(4.1), Inches(1.8)),
        ("Execution Engine", "FastAPI Core & DB State\n(Executes Task ID &\nCapacity Drop)", Inches(8.3), Inches(4.7), Inches(4.1), Inches(1.8))
    ]
    for stage, desc, left, top, width, height in blocks:
        box = slide4.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = stage
        p1.font.bold = True
        p1.font.size = Pt(14)
        p1.font.color.rgb = RGBColor(37, 99, 235)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = RGBColor(30, 41, 59)

    # --- SLIDE 5: Solution Formulation – Explanation ---
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide5, RGBColor(255, 255, 255))
    add_title(slide5, "Solution Formulation – Technical Explanation")

    cbox5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf5 = cbox5.text_frame
    tf5.word_wrap = True

    sol_points = [
        ("Decoupled Hybrid Architecture: ", "Isolates probabilistic generation (Llama 3) from factual calculations. All date offsets, calendar overlaps, and capacity hours are computed deterministically in Python."),
        ("Strict Schema Enforcement: ", "Pydantic models act as runtime security boundaries. Unstructured outputs are coerced into strictly validated data structures via the Instructor library."),
        ("Factual Grounding Loop: ", "The LLM reasoning node receives structured candidate dossiers (exact capacity hours, matching skills, missing proficiencies) rather than querying the database directly, avoiding hallucinations."),
        ("Non-Mutating Sandbox: ", "The What-If Simulator clones in-memory state to model secondary impacts of sudden leaves or deadline shifts without altering production tasks.")
    ]
    for i, (title, desc) in enumerate(sol_points):
        p = tf5.paragraphs[0] if i == 0 else tf5.add_paragraph()
        p.space_after = Pt(14)
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(17)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(17)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # --- SLIDE 6: Traditional vs GenAI Workflow ---
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide6, RGBColor(255, 255, 255))
    add_title(slide6, "Traditional Workflow vs GenAI Workflow")

    table_shape = slide6.shapes.add_table(6, 3, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = table_shape.table

    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(4.45)
    table.columns[2].width = Inches(4.45)

    headers = ["Lifecycle Stage", "Traditional Workflow", "GenAI Workflow (AI Delivery Manager)"]
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(37, 99, 235)
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(255, 255, 255)

    rows = [
        ("Task Creation", "Manual input across 10+ form fields (Jira/Planner).", "Natural language parsing into validated JSON models."),
        ("Capacity Verification", "Manual calendar inspections and static spreadsheets.", "Deterministic engine computes net remaining hours."),
        ("Candidate Allocation", "Intuitive selection or rigid static performance metrics.", "Algorithmic skill match paired with LLM trade-off analysis."),
        ("Absence Management", "Discovered reactively after project delays occur.", "Conversational agent analyzes blast radius instantly."),
        ("Execution Authority", "Direct managerial assignment.", "AI recommends → Human approves → System executes (HITL).")
    ]
    for row_idx, data in enumerate(rows, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.color.rgb = RGBColor(30, 41, 59)

    # --- SLIDE 7: Usage of Gen AI Tools ---
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide7, RGBColor(255, 255, 255))
    add_title(slide7, "Usage of GenAI Tools")

    cbox7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf7 = cbox7.text_frame
    tf7.word_wrap = True

    tools = [
        ("Ollama Runtime: ", "Hosts open-weight models locally, ensuring enterprise data privacy and eliminating cloud inference latency and costs."),
        ("Meta Llama 3 (8B Instruct): ", "Serves as the core language intelligence layer for task extraction, conversational agent analysis, and recommendation reasoning."),
        ("Instructor Library: ", "Enforces strict JSON schema validation and retry mechanisms directly on Llama 3 outputs via OpenAI-compatible endpoints."),
        ("FastAPI & Pydantic: ", "Provides the backend orchestration, request routing, and deterministic data contract validation.")
    ]
    for i, (tool, desc) in enumerate(tools):
        p = tf7.paragraphs[0] if i == 0 else tf7.add_paragraph()
        p.space_after = Pt(16)
        r1 = p.add_run()
        r1.text = "• " + tool
        r1.font.bold = True
        r1.font.size = Pt(18)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(18)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # --- SLIDE 8: Challenge Identification & Mitigation ---
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide8, RGBColor(255, 255, 255))
    add_title(slide8, "Challenge Identification & Mitigation")

    cbox8 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf8 = cbox8.text_frame
    tf8.word_wrap = True

    challenges = [
        ("Arithmetic & Schedule Hallucinations: ", "LLMs fail at calculating working hours and time deltas.\n→ Mitigation: Abstracted math into pure Python engines; LLM only reasons on verified integers."),
        ("JSON Malformation in Local Models: ", "Smaller 8B models can produce invalid JSON that breaks API routes.\n→ Mitigation: Integrated the Instructor framework to enforce automatic schema retry loops."),
        ("Unauthorized Autonomous State Mutation: ", "AI directly executing business assignments poses operational risks.\n→ Mitigation: Architected an immutable Human-in-the-Loop approval gate before state commits."),
        ("Data Privacy & HR Sensitivity: ", "Sending workforce schedules and private calendar items to third-party cloud APIs poses security risks.\n→ Mitigation: Fully local execution using Ollama on dedicated edge hardware.")
    ]
    for i, (chal, desc) in enumerate(challenges):
        p = tf8.paragraphs[0] if i == 0 else tf8.add_paragraph()
        p.space_after = Pt(14)
        r1 = p.add_run()
        r1.text = "• " + chal
        r1.font.bold = True
        r1.font.size = Pt(17)
        r1.font.color.rgb = RGBColor(220, 38, 38)
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(16)
        r2.font.color.rgb = RGBColor(51, 65, 85)

    # --- SLIDE 9: Evaluation Metrics ---
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide9, RGBColor(255, 255, 255))
    add_title(slide9, "Evaluation Metrics & Performance Validation")

    cbox9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf9 = cbox9.text_frame
    tf9.word_wrap = True

    metrics = [
        ("Schema Adherence Rate (>95%): ", "Measures how consistently local Llama 3 generates valid JSON conforming to Pydantic schemas without exhaustion of retries."),
        ("Mathematical Determinism (100%): ", "Zero margin of error on available hours and calendar math by using deterministic execution."),
        ("HITL Acceptance Rate: ", "Quantifies the percentage of AI-generated allocation suggestions approved by the human manager without modification."),
        ("Inference & Execution Latency: ", "Tracks end-to-end response times across task parsing, candidate scoring, and reasoning synthesis on consumer-grade hardware (NVIDIA RTX 3050 Ti).")
    ]
    for i, (metric, desc) in enumerate(metrics):
        p = tf9.paragraphs[0] if i == 0 else tf9.add_paragraph()
        p.space_after = Pt(16)
        r1 = p.add_run()
        r1.text = "• " + metric
        r1.font.bold = True
        r1.font.size = Pt(18)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(18)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # --- SLIDE 10: Innovation & Customization ---
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide10, RGBColor(255, 255, 255))
    add_title(slide10, "Customization & Innovation Component(s)")

    cbox10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf10 = cbox10.text_frame
    tf10.word_wrap = True

    innovations = [
        ("Local Edge Deployment: ", "Zero-cost, enterprise-compliant local architecture running on edge GPUs, safeguarding organizational data from external cloud leakage."),
        ("Conversational Blast Radius Analyzer: ", "Empowers employees to communicate in unstructured natural language while dynamically mapping project downstream risks."),
        ("Non-Destructive Digital Twin Simulation: ", "Enables risk-free workforce schedule forecasting and rebalancing without modifying live task states.")
    ]
    for i, (title, desc) in enumerate(innovations):
        p = tf10.paragraphs[0] if i == 0 else tf10.add_paragraph()
        p.space_after = Pt(18)
        r1 = p.add_run()
        r1.text = "• " + title
        r1.font.bold = True
        r1.font.size = Pt(18)
        r1.font.color.rgb = RGBColor(37, 99, 235)
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(18)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # --- SLIDE 11: Thank You ---
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide11, RGBColor(248, 250, 252))

    tbox11 = slide11.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.3), Inches(3.0))
    tf11 = tbox11.text_frame
    p_end = tf11.paragraphs[0]
    p_end.text = "Thank You!"
    p_end.font.size = Pt(48)
    p_end.font.bold = True
    p_end.font.color.rgb = RGBColor(37, 99, 235)
    p_end.alignment = PP_ALIGN.CENTER

    p_qa = tf11.add_paragraph()
    p_qa.text = "Questions & Demonstration"
    p_qa.font.size = Pt(22)
    p_qa.font.color.rgb = RGBColor(100, 116, 139)
    p_qa.alignment = PP_ALIGN.CENTER

    output_filename = "25ARE471_Introd_to_GenAI_CB.EN.U4ARE23023.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully created: {output_filename}")

if __name__ == "__main__":
    create_presentation()